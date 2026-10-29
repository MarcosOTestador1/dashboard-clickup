# -*- coding: utf-8 -*-
"""Backend Flask do Dashboard ClickUp: filtro por data, retries e cache de contingência."""
import logging
import os
import threading
import time
from datetime import datetime, timedelta, timezone
from io import BytesIO

import requests
from flask import Flask, jsonify, render_template, request, send_file

app = Flask(__name__, template_folder="templates")

API_BASE = "https://api.clickup.com/api/v2"
DIAS_MINIMOS = 100
CONNECT_TIMEOUT = 5
READ_TIMEOUT = 12
MAX_RETRIES = 2
CACHE_TTL_SECONDS = 300
MAX_PAGES = 10  # limite temporário para evitar buscas intermináveis
MAX_FETCH_SECONDS = 75  # orçamento máximo aproximado para a paginação

STATUS_INATIVOS = {
    "cancelado", "cancelada", "cancelados", "canceladas",
    "cancelled", "canceled", "finalizado", "finalizada",
    "finalizados", "finalizadas", "closed", "done",
    "concluido", "concluida", "concluídos", "concluídas",
    "complete", "completed",
}

_cache_lock = threading.Lock()
_refresh_lock = threading.Lock()
_cache = {"at": 0.0, "tasks": None, "updated_at": None}


def _safe_json(response):
    """Valida HTTP e JSON, evitando tratar HTML de erro como JSON."""
    response.raise_for_status()
    try:
        payload = response.json()
    except (requests.exceptions.JSONDecodeError, ValueError) as exc:
        content_type = response.headers.get("Content-Type", "")
        app.logger.warning(
            "Resposta não-JSON do ClickUp (HTTP %s, Content-Type %s).",
            response.status_code, content_type
        )
        raise ValueError("O ClickUp retornou uma resposta que não é JSON válido.") from exc

    if not isinstance(payload, dict):
        raise ValueError("Resposta inesperada da API do ClickUp.")
    return payload


def _get_json(session, url, params=None):
    """Repete falhas transitórias de rede, HTTP 429 e HTTP 5xx."""
    last_error = None

    for attempt in range(MAX_RETRIES):
        try:
            response = session.get(
                url,
                params=params,
                timeout=(CONNECT_TIMEOUT, READ_TIMEOUT),
            )

            transient = response.status_code == 429 or 500 <= response.status_code <= 599
            if transient and attempt < MAX_RETRIES - 1:
                retry_after = response.headers.get("Retry-After")
                try:
                    delay = min(8.0, max(1.0, float(retry_after))) if retry_after else 2 ** attempt
                except (TypeError, ValueError):
                    delay = 2 ** attempt
                app.logger.warning(
                    "ClickUp respondeu HTTP %s; nova tentativa em %.1fs (%s/%s).",
                    response.status_code, delay, attempt + 1, MAX_RETRIES
                )
                time.sleep(delay)
                continue

            return _safe_json(response)

        except (requests.Timeout, requests.ConnectionError) as exc:
            last_error = exc
            if attempt >= MAX_RETRIES - 1:
                raise
            delay = 2 ** attempt
            app.logger.warning(
                "Falha temporária de rede com ClickUp (%s); nova tentativa em %ss (%s/%s).",
                type(exc).__name__, delay, attempt + 1, MAX_RETRIES
            )
            time.sleep(delay)

    if last_error:
        raise last_error
    raise requests.RequestException("Não foi possível consultar a API do ClickUp.")


def _space_name(task, spaces_by_id):
    space = task.get("space") or {}
    if isinstance(space, dict):
        name = space.get("name")
        space_id = space.get("id")
        if name:
            return str(name)
    else:
        space_id = None

    space_id = space_id or task.get("space_id")
    if space_id is not None and str(space_id) in spaces_by_id:
        return spaces_by_id[str(space_id)]

    list_obj = task.get("list") or {}
    if isinstance(list_obj, dict) and list_obj.get("space"):
        nested = list_obj["space"]
        if isinstance(nested, dict) and nested.get("name"):
            return str(nested["name"])
        if isinstance(nested, str) and nested in spaces_by_id:
            return spaces_by_id[nested]

    if isinstance(list_obj, dict) and list_obj.get("name"):
        return f"Lista: {list_obj['name']}"
    return "Espaço não identificado"


def _is_inactive_status(status):
    normalized = (status or "").lower().strip()
    return normalized in STATUS_INATIVOS or any(
        term in normalized
        for term in (
            "cancelad", "cancelled", "canceled", "finalizad",
            "conclu", "complet", "closed", "done",
        )
    )


def buscar_tarefas_clickup():
    """Busca tarefas abertas criadas há mais de DIAS_MINIMOS dias."""
    token = (os.getenv("CLICKUP_API_TOKEN") or "").strip()
    if not token:
        raise RuntimeError(
            "O serviço ainda não foi configurado com CLICKUP_API_TOKEN no Render."
        )

    session = requests.Session()
    session.headers.update({
        "Authorization": token,
        "Accept": "application/json",
    })

    try:
        teams = _get_json(session, f"{API_BASE}/team").get("teams", [])
        if not teams:
            raise RuntimeError("Nenhum workspace do ClickUp foi encontrado para este token.")

        team_id = (os.getenv("CLICKUP_TEAM_ID") or "").strip()
        team = next((t for t in teams if str(t.get("id")) == team_id), None) if team_id else teams[0]
        if not team:
            raise RuntimeError(
                "CLICKUP_TEAM_ID não corresponde a um workspace disponível para o token."
            )

        spaces_by_id = {}
        try:
            spaces_payload = _get_json(
                session,
                f"{API_BASE}/team/{team['id']}/space",
                {"archived": "false"},
            )
            spaces_by_id = {
                str(space["id"]): str(space["name"])
                for space in spaces_payload.get("spaces", [])
                if space.get("id") is not None and space.get("name")
            }
        except (requests.RequestException, ValueError) as exc:
            # Os nomes dos espaços são complementares; a busca de tarefas pode continuar.
            app.logger.warning(
                "Não foi possível carregar a lista de Spaces: %s", type(exc).__name__
            )

        now = datetime.now(timezone.utc)
        cutoff_ms = int((now - timedelta(days=DIAS_MINIMOS)).timestamp() * 1000)
        critical = []
        seen_ids = set()
        fetch_started = time.monotonic()

        for page in range(MAX_PAGES):
            if time.monotonic() - fetch_started >= MAX_FETCH_SECONDS:
                app.logger.warning(
                    "Busca interrompida após %ss para evitar timeout; resultados podem estar incompletos.",
                    MAX_FETCH_SECONDS,
                )
                break

            try:
                payload = _get_json(
                    session,
                    f"{API_BASE}/team/{team['id']}/task",
                    {
                        "page": page,
                        "subtasks": "true",
                        "date_created_lt": cutoff_ms,
                        "include_closed": "false",
                    },
                )
            except (requests.RequestException, ValueError) as exc:
                # Se já recebemos páginas válidas, preserva os dados coletados
                # em vez de descartar tudo por uma falha numa página posterior.
                if critical:
                    app.logger.warning(
                        "Falha ao consultar página %s do ClickUp (%s); retornando resultados parciais.",
                        page, type(exc).__name__,
                    )
                    break
                raise

            tasks = payload.get("tasks", [])
            if not tasks:
                break

            for task in tasks:
                task_id = str(task.get("id") or task.get("custom_id") or "")
                # Evita duplicação caso a API devolva tarefas repetidas entre páginas.
                if task_id and task_id in seen_ids:
                    continue
                if task_id:
                    seen_ids.add(task_id)

                status_obj = task.get("status") or {}
                status = (
                    status_obj.get("status", "")
                    if isinstance(status_obj, dict)
                    else str(status_obj)
                )
                if _is_inactive_status(status):
                    continue

                created_ms = task.get("date_created")
                if created_ms in (None, ""):
                    continue
                try:
                    created = datetime.fromtimestamp(
                        int(created_ms) / 1000, tz=timezone.utc
                    )
                except (TypeError, ValueError, OSError, OverflowError):
                    continue

                days = max(0, (now - created).days)
                if days <= DIAS_MINIMOS:
                    continue

                assignees = task.get("assignees") or []
                responsible_names = []
                for assignee in assignees:
                    if isinstance(assignee, dict):
                        name = (
                            assignee.get("username")
                            or assignee.get("email")
                            or assignee.get("initials")
                        )
                        if name:
                            responsible_names.append(str(name))
                responsible = ", ".join(responsible_names) or "Não atribuído"

                critical.append({
                    "id": str(task.get("custom_id") or task.get("id") or ""),
                    "name": str(task.get("name") or "Sem título"),
                    "status": status or "Sem status",
                    "dias": days,
                    "criacao": created.strftime("%Y-%m-%d"),
                    "responsavel": responsible,
                    "espaco": _space_name(task, spaces_by_id),
                    "url": str(task.get("url") or "#"),
                })
        else:
            app.logger.warning(
                "A consulta atingiu o limite de %s páginas; resultados podem estar incompletos.",
                MAX_PAGES
            )

        return critical
    finally:
        session.close()


def obter_tarefas_cache(force_refresh=False):
    """
    Retorna cache recente quando disponível.
    Uma trava separada evita consultas simultâneas; o lock do cache nunca fica
    preso durante a chamada de rede. Em falha, usa a última consulta bem-sucedida.
    """
    now = time.time()
    with _cache_lock:
        has_cache = _cache["tasks"] is not None
        fresh = has_cache and (now - _cache["at"] < CACHE_TTL_SECONDS)
        if not force_refresh and fresh:
            return list(_cache["tasks"]), _cache["updated_at"], True

    # Só uma thread atualiza por vez. As demais podem usar o cache existente.
    if not _refresh_lock.acquire(blocking=False):
        with _cache_lock:
            if _cache["tasks"] is not None:
                return list(_cache["tasks"]), _cache["updated_at"], True
        # Sem cache ainda: aguarda a atualização que já está em andamento.
        with _refresh_lock:
            with _cache_lock:
                if _cache["tasks"] is not None:
                    return list(_cache["tasks"]), _cache["updated_at"], True
            # Se a primeira atualização falhou sem cache, tenta uma vez nesta thread.
            return _buscar_e_atualizar_cache()

    try:
        return _buscar_e_atualizar_cache()
    finally:
        _refresh_lock.release()


def _buscar_e_atualizar_cache():
    try:
        tasks = buscar_tarefas_clickup()
        updated_at = datetime.now(timezone.utc).isoformat()
        with _cache_lock:
            _cache.update(at=time.time(), tasks=tasks, updated_at=updated_at)
        return list(tasks), updated_at, False
    except Exception:
        with _cache_lock:
            has_cache = _cache["tasks"] is not None
            cached_tasks = list(_cache["tasks"]) if has_cache else None
            cached_updated_at = _cache["updated_at"]

        if has_cache:
            app.logger.exception(
                "Atualização do ClickUp falhou; retornando o último cache bem-sucedido (%s).",
                cached_updated_at,
            )
            return cached_tasks, cached_updated_at, True
        raise


@app.after_request
def security_headers(response):
    response.headers["X-Content-Type-Options"] = "nosniff"
    response.headers["X-Frame-Options"] = "SAMEORIGIN"
    response.headers["Referrer-Policy"] = "strict-origin-when-cross-origin"
    response.headers["Cache-Control"] = "no-store"
    return response


@app.get("/")
def index():
    return render_template("index.html")


@app.get("/api/tasks")
def api_tasks():
    try:
        force_refresh = request.args.get("refresh") == "1"
        tasks, updated_at, cached = obter_tarefas_cache(force_refresh)
        return jsonify({
            "tasks": tasks,
            "updated_at": updated_at,
            "cached": cached,
        }), 200

    except requests.HTTPError as exc:
        code = exc.response.status_code if exc.response is not None else "desconhecido"
        app.logger.warning("ClickUp API retornou HTTP %s", code)
        return jsonify({
            "error": "Não foi possível consultar o ClickUp. Verifique o token e as permissões configuradas no Render."
        }), 502

    except (requests.RequestException, ValueError) as exc:
        app.logger.warning("Falha de comunicação com ClickUp: %s", type(exc).__name__)
        return jsonify({
            "error": "Falha de comunicação com o ClickUp. Tente novamente em alguns instantes."
        }), 502

    except RuntimeError as exc:
        app.logger.warning("Configuração ClickUp: %s", exc)
        return jsonify({"error": str(exc)}), 503

    except Exception:
        app.logger.exception("Erro inesperado ao carregar tarefas do ClickUp.")
        return jsonify({
            "error": "Erro interno ao carregar as tarefas. Consulte os logs do servidor."
        }), 500


@app.post("/api/export.xlsx")
def export_xlsx():
    """Exporta as tarefas enviadas pela tabela filtrada do dashboard."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        return jsonify({
            "error": "Exportação XLSX indisponível: dependência openpyxl não instalada."
        }), 503

    payload = request.get_json(silent=True) or {}
    rows = payload.get("tasks")
    if not isinstance(rows, list):
        return jsonify({
            "error": "Não foi possível identificar as tarefas para exportação."
        }), 400
    if len(rows) > 20000:
        return jsonify({
            "error": "A seleção excede o limite de 20.000 tarefas por exportação."
        }), 413

    headers = [
        "ID", "Nome da tarefa", "Status", "Dias em aberto", "Responsável",
        "Espaço do ClickUp", "Data de criação", "Link ClickUp",
    ]
    fields = ["id", "name", "status", "dias", "responsavel", "espaco", "criacao", "url"]

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Detalhamento das Tarefas"
    sheet.append(headers)

    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="334155")
        cell.alignment = Alignment(vertical="center", wrap_text=True)

    def safe_cell(value):
        if value is None:
            return ""
        if isinstance(value, (int, float)) and not isinstance(value, bool):
            return value
        value = str(value)
        # Mitiga fórmulas maliciosas em valores exportados.
        return "'" + value if value.startswith(("=", "+", "-", "@")) else value

    for item in rows:
        if not isinstance(item, dict):
            continue
        values = [safe_cell(item.get(field, "")) for field in fields]
        url = str(item.get("url") or "")
        if not (
            url.startswith("https://app.clickup.com/")
            or url.startswith("https://sharing.clickup.com/")
        ):
            values[-1] = ""

        sheet.append(values)
        link = sheet.cell(row=sheet.max_row, column=8)
        if values[-1]:
            link.hyperlink = values[-1]
            link.style = "Hyperlink"

        days_cell = sheet.cell(row=sheet.max_row, column=4)
        if isinstance(days_cell.value, (int, float)):
            days_cell.number_format = '0 "dias"'

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.row_dimensions[1].height = 30
    for index, width in enumerate([18, 55, 22, 18, 32, 32, 18, 48], start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=(cell.column == 2),
            )

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    filename = f"detalhamento_tarefas_clickup_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return send_file(
        output,
        as_attachment=True,
        download_name=filename,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@app.get("/health")
def health():
    return jsonify({"status": "ok"})


if __name__ == "__main__":
    logging.basicConfig(level=os.getenv("LOG_LEVEL", "INFO"))
    app.run(
        host="0.0.0.0",
        port=int(os.getenv("PORT", "5000")),
        debug=False,
    )
